"""Load a Bloomberg MODL .xlsx export into the valuation.db SQLite schema.

A raw Bloomberg MODL export normally ships ONE sheet: "Multiple Periods"
(annual actuals + Street-consensus estimates, one column per fiscal year,
rows identified by a Bloomberg field code in column B). A hand-built "DCF"
tab (Bear/Base/Bull assumptions, terminal exit multiple) and "WACC" tab
(CAPM inputs) are NOT something Bloomberg exports — the AppLovin workbook
this project started from had them because a human built them by hand.

So this loader supports two modes:

  - Mode A ("DCF"/"WACC" tabs present): extract the human-built assumptions
    verbatim, exactly as before. This is how AppLovin's validated $388.54
    base-case price was produced, and this path is unchanged.
  - Mode B ("Multiple Periods" only, the normal case): DERIVE everything
    ourselves. Scenario assumptions (revenue growth, margins, tax rate,
    D&A%/CapEx%/NWC%) come from `derive_scenario_assumptions()`, applied
    formulaically to the historicals. The terminal EV/EBITDA exit multiple
    defaults to the company's own current market-implied multiple. WACC's
    CAPM inputs (risk-free rate, beta, equity risk premium, stock price)
    are NOT in any Bloomberg financials export — they're market data, not
    company data — so they must be supplied explicitly via `market_data`;
    this loader never fabricates them.

Bloomberg's field codes for a given concept (e.g. "pre-tax income") are
not stable across industry templates — the same code can mean different
things for an operating company vs. an alternative asset manager, and a
concept can use different codes entirely. `historicals` is therefore
built from a per-concept alias list, trying the most standard/universal
Bloomberg code first.

Usage:
    python data/load_from_modl.py <path-to-modl.xlsx> <path-to-valuation.db>
    python data/load_from_modl.py <path-to-modl.xlsx> <path-to-valuation.db> \
        --risk-free-rate 0.0477 --beta 1.75 --erp 0.0445 --stock-price 136.82
"""

from __future__ import annotations

import argparse
import re
import sqlite3
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

# Columns C..O on a hand-built DCF sheet hold FY2026E..FY2038E (13 years)
# for every assumption block (Base, Bear, Bull) and for the FCF build.
# Mode B's derived assumptions target the same 13-year window so both
# modes plug into dcf_engine.py identically.
YEAR_COLS = ["C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O"]
FIRST_FORECAST_YEAR = 2026
TOTAL_FORECAST_YEARS = len(YEAR_COLS)  # 13


# Bloomberg field-code alias chains per historicals concept, most standard/
# universal code first. Required: the loader raises if NONE of a concept's
# aliases are found anywhere in the sheet — that concept is load-bearing
# for dcf_engine.py's math and a silent None would only surface later as a
# confusing arithmetic error.
REQUIRED_FIELD_ALIASES: dict[str, list[str]] = {
    "revenue": ["IS_COMP_SALES", "SALES_REV_TURN"],
    "ebit": ["IS_COMPARABLE_EBIT"],
    "ebitda_adjusted": ["IS_COMPARABLE_EBITDA"],
    "da": ["CF_DEPR_AMORT", "CB_IS_DEPRECIATION_AMORT_EXP"],
    "pretax_income": ["PRETAX_INC", "IS_COMP_PTP_EX_STK_BASED_COMP"],
    "tax_expense": ["IS_INC_TAX_EXP"],
    "net_income": ["IS_COMP_NET_INCOME_GAAP"],
    "eps_diluted_adjusted": ["IS_COMP_EPS_ADJUSTED_OLD", "IS_COMP_EPS_GAAP"],
    "long_term_debt": ["CB_BS_LT_BORROWING", "BS_LONG_TERM_BORROWINGS"],
    "capex": ["CB_CF_PURCHASES_OF_PPE", "HEADLINE_CAPEX"],
    "shares_diluted": ["IS_SH_FOR_DILUTED_EPS"],
}

# Optional concepts: default rather than raise when no alias is found.
# net_debt falls back to long_term_debt - cash; nwc_change defaults to 0
# (a documented simplification — some business models, e.g. alternative
# asset managers, genuinely don't have a meaningful working-capital cycle).
# interest_expense is optional (rather than required) because its Bloomberg
# field code varies more than any other concept here across industry
# templates -- when none of these aliases match, it defaults to 0.0 (cost
# of debt then reads as 0% rather than crashing the whole upload; see
# load_historicals()) and app.py surfaces that explicitly.
OPTIONAL_FIELD_ALIASES: dict[str, list[str]] = {
    "net_debt": ["NET_DEBT"],
    "nwc_change": ["CF_CHNG_NON_CASH_WORK_CAP"],
    "cash": ["BS_CASH_CASH_EQUIVALENTS_AND_STI", "BS_CASH_NEAR_CASH_ITEM"],
    "interest_expense": [
        "IS_NET_INTEREST_EXPENSE", "CB_IS_INTEREST_EXPENSE", "IS_INT_EXPENSES",
        "IS_INTEREST_EXPENSE", "IS_NET_INTEREST_INC_EXP", "FIN_NET_INT_EXP",
        "IS_TOT_INT_EXP", "NET_INT_EXP_GAAP", "IS_INTEREST_EXPENSE_NET",
        "CB_IS_NET_INTEREST_EXPENSE", "IS_INTEREST_INC_EXP_NET",
    ],
}

# Long-run nominal growth used to fade Mode B's revenue growth down to
# after the Street-consensus window runs out (roughly long-run nominal
# GDP growth — inflation + real growth — a standard terminal-growth proxy).
TERMINAL_GROWTH_RATE = 0.03

# Mode B bear/bull spread applied to the derived base case: a documented,
# transparent rule (not extracted data) for producing a sensitivity range
# when there's no hand-built scenario table to read from.
BEAR_GROWTH_MULTIPLIER = 0.7
BULL_GROWTH_MULTIPLIER = 1.3
BEAR_MARGIN_DELTA = -0.03
BULL_MARGIN_DELTA = 0.02
BEAR_MULTIPLE_MULTIPLIER = 0.85
BULL_MULTIPLE_MULTIPLIER = 1.15

# (scenario, header row of the assumption block on a hand-built DCF sheet).
# Each block is: header (years), +1 revenue growth, +2 EBIT margin, +3 tax
# rate, +4 D&A %, +5 CapEx %, +6 NWC %, +7 terminal exit multiple (col C only).
SCENARIO_BLOCKS = {
    "base": 16,
    "bear": 26,
    "bull": 36,
}

# DCF sheet "TRADING MULTIPLES" block: row -> (metric name, {column: fiscal_year}).
TRADING_COMPS_ROWS = {
    130: "PE",
    131: "EV_EBITDA",
    132: "FCF_YIELD",
}
TRADING_COMPS_COLS = {"B": 2027, "C": 2030}

REQUIRED_SHEETS_MODE_A = ("Multiple Periods", "DCF", "WACC")
REQUIRED_MARKET_DATA_KEYS = ("risk_free_rate", "beta", "equity_risk_premium", "stock_price")


def _numeric_or_none(value):
    """Bloomberg sometimes leaves a gap in an otherwise-numeric estimate
    series as an empty string rather than a blank cell -- normalize that
    to None so it doesn't get stored as text in a REAL column (SQLite
    doesn't enforce column types) and only blow up arithmetic later, far
    from the actual cause."""
    if isinstance(value, str) and not value.strip():
        return None
    return value


def find_all_field_rows(ws) -> dict[str, int]:
    """Map every Bloomberg field code appearing in column B to its first row."""
    rows: dict[str, int] = {}
    for (cell,) in ws.iter_rows(min_col=2, max_col=2):
        if cell.value and cell.value not in rows:
            rows[cell.value] = cell.row
    return rows


def find_year_columns(ws, header_row: int = 3) -> list[str]:
    """Year columns on the historicals sheet, detected from the period-label
    row ('2026 A (Fwd)', '2025 A (Rep)', ...) instead of assumed to sit in a
    fixed E:N range. A field-code MODL export happens to start at column E
    (10 years); a raw export saved straight from Bloomberg's on-screen grid
    (no field-code column at all) starts at column B instead and can span a
    different number of years -- this reads whichever columns the sheet
    actually has."""
    return [cell.column_letter for cell in ws[header_row] if cell.value and re.match(r"^\d{4}\s", str(cell.value))]


def resolve_concept(all_rows: dict[str, int], aliases: list[str]) -> int | None:
    """First alias code (in priority order) that has a row in this sheet."""
    for code in aliases:
        if code in all_rows:
            return all_rows[code]
    return None


def period_type_for_column(ws, col: str) -> str:
    """Row 3 tags each year column '... (Fwd)' (estimate) or '... (Rep)' (actual)."""
    label = ws[f"{col}3"].value or ""
    if "Fwd" in label:
        return "estimate"
    if "Rep" in label:
        return "actual"
    raise ValueError(f"Cannot determine period type for column {col} from label {label!r}")


def load_historicals(ws) -> list[dict]:
    all_rows = find_all_field_rows(ws)

    required_rows = {concept: resolve_concept(all_rows, aliases) for concept, aliases in REQUIRED_FIELD_ALIASES.items()}
    missing = [concept for concept, row in required_rows.items() if row is None]
    if missing:
        raise ValueError(
            f"Could not find any known Bloomberg field code for {missing} in 'Multiple "
            f"Periods'. Tried: {[REQUIRED_FIELD_ALIASES[c] for c in missing]}."
        )
    optional_rows = {concept: resolve_concept(all_rows, aliases) for concept, aliases in OPTIONAL_FIELD_ALIASES.items()}

    records = []
    for col in find_year_columns(ws):
        year = ws[f"{col}4"].value.year
        record = {"fiscal_year": year, "period_type": period_type_for_column(ws, col)}
        for concept, row in required_rows.items():
            record[concept] = _numeric_or_none(ws[f"{col}{row}"].value)
        for concept, row in optional_rows.items():
            record[concept] = _numeric_or_none(ws[f"{col}{row}"].value) if row else None

        if record["net_debt"] is None:
            record["net_debt"] = (record["long_term_debt"] or 0) - (record["cash"] or 0)
        record.pop("cash", None)
        if record["nwc_change"] is None:
            record["nwc_change"] = 0.0
        if record["interest_expense"] is None:
            record["interest_expense"] = 0.0

        records.append(record)
    return records


def load_wacc_inputs(wacc_ws) -> dict:
    return {
        "risk_free_rate": wacc_ws["B5"].value,
        "beta": wacc_ws["B6"].value,
        "equity_risk_premium": wacc_ws["B7"].value,
        "stock_price": wacc_ws["B18"].value,
        "shares_outstanding": wacc_ws["B19"].value,
        "source": "Aba WACC do arquivo original (células B5, B6, B7, B18)",
    }


def load_scenario_block(dcf_ws, header_row: int) -> tuple[list[dict], float]:
    growth_row, margin_row, tax_row = header_row + 1, header_row + 2, header_row + 3
    da_row, capex_row, nwc_row = header_row + 4, header_row + 5, header_row + 6
    exit_multiple_row = header_row + 7

    assumptions = []
    for col, year in zip(YEAR_COLS, range(FIRST_FORECAST_YEAR, FIRST_FORECAST_YEAR + len(YEAR_COLS))):
        assumptions.append({
            "fiscal_year": year,
            "revenue_growth": dcf_ws[f"{col}{growth_row}"].value,
            "ebit_margin": dcf_ws[f"{col}{margin_row}"].value,
            "tax_rate": dcf_ws[f"{col}{tax_row}"].value,
            "da_pct_revenue": dcf_ws[f"{col}{da_row}"].value,
            "capex_pct_revenue": dcf_ws[f"{col}{capex_row}"].value,
            "nwc_pct_delta_revenue": dcf_ws[f"{col}{nwc_row}"].value,
        })
    exit_multiple = dcf_ws[f"C{exit_multiple_row}"].value
    return assumptions, exit_multiple


def load_trading_comps(dcf_ws) -> list[dict]:
    comps = []
    for row, metric in TRADING_COMPS_ROWS.items():
        for col, year in TRADING_COMPS_COLS.items():
            comps.append({"metric": metric, "fiscal_year": year, "value": dcf_ws[f"{col}{row}"].value})
    # PEG (row 133) is a single value anchored to the block's first column (FY2027E).
    comps.append({"metric": "PEG", "fiscal_year": TRADING_COMPS_COLS["B"], "value": dcf_ws["B133"].value})
    return comps


def parse_ticker(multiple_periods_ws) -> str:
    header = multiple_periods_ws["A2"].value or ""
    match = re.match(r"^(\S+\s+\S+)\s+Equity", header)
    if not match:
        raise ValueError(f"Could not parse ticker from 'Multiple Periods'!A2: {header!r}")
    return match.group(1)


def parse_company_name_from_title(title: str) -> str:
    match = re.match(r"^(.+?)\s*[-(]", title)
    return match.group(1).strip() if match else title.strip()


def parse_as_of_date(header: str) -> str | None:
    match = re.search(r"As of:\s*([\d/]+)", header)
    if not match:
        return None
    return datetime.strptime(match.group(1), "%m/%d/%Y").date().isoformat()


def find_multiple_periods_sheet(wb):
    """The historicals/consensus sheet, found by its title text (cell A1
    ends in "... (Multiple Periods)") rather than its tab name. A workbook
    saved straight from Bloomberg with no manual renaming can leave the tab
    itself named "Planilha1"/"Sheet1" (whatever the spreadsheet app's
    default is) while the sheet's own title still reads correctly. Falls
    back to a tab literally named "Multiple Periods" for workbooks with no
    title in A1."""
    for ws in wb.worksheets:
        title = ws["A1"].value
        if title and "multiple periods" in str(title).lower():
            return ws
    return wb["Multiple Periods"] if "Multiple Periods" in wb.sheetnames else None


# --------------------------------------- label-only export (no field codes) --
# A workbook saved straight from Bloomberg's on-screen "Company Financial
# (Multiple Periods)" grid -- rather than exported via the field-code-driven
# MODL template every other path in this module assumes -- has no BDH-style
# field-code column: column B holds the first year's data, not a code, and
# year columns start wherever the screen put them (column B, not E). Labels
# repeat across the sheet (e.g. "Revenue" and "Net Income" both appear under
# segment/business breakdowns AND under the actual financial statements), so
# each concept is resolved within a named section -- "Income Statement",
# "Condensed Balance Sheet", "Condensed Cash Flow Statement" (Bloomberg's
# own, standard headers for this screen) -- at a specific indent depth,
# rather than by a bare text search.

LABEL_SECTIONS = ("Income Statement", "Condensed Balance Sheet", "Condensed Cash Flow Statement")

# concept -> (section, exact line-item label, indent depth within that section)
REQUIRED_LABEL_MAP: dict[str, tuple[str, str, int]] = {
    "revenue": ("Income Statement", "Total Revenue", 1),
    "ebit": ("Income Statement", "Operating Income", 1),
    "ebitda_adjusted": ("Income Statement", "EBITDA", 2),
    "da": ("Income Statement", "Depreciation & Amortization", 1),
    # Gross interest expense, not the net-of-interest-income figure: a
    # net-cash company (e.g. Alphabet) has negative net interest, which
    # would otherwise make dcf_engine.get_wacc()'s pretax cost of debt
    # negative -- a company's cost of its own debt shouldn't depend on how
    # much interest income its cash balance happens to throw off.
    "interest_expense": ("Income Statement", "Interest Expense", 2),
    "pretax_income": ("Income Statement", "Pre-Tax Income", 1),
    "tax_expense": ("Income Statement", "Income Tax Expense", 1),
    "net_income": ("Income Statement", "Net Income", 1),
    "eps_diluted_adjusted": ("Income Statement", "Diluted EPS", 2),
    "long_term_debt": ("Condensed Balance Sheet", "Long-Term Debt", 3),
    "capex": ("Condensed Cash Flow Statement", "Capital Expenditures", 2),
    "shares_diluted": ("Income Statement", "Diluted Weighted Avg. Shares", 1),
}
OPTIONAL_LABEL_MAP: dict[str, tuple[str, str, int]] = {
    "cash": ("Condensed Balance Sheet", "Cash Cash Equivalents and Short term Investments", 3),
    "nwc_change": ("Condensed Cash Flow Statement", "Changes in Working Capital", 2),
}


def _row_label_depth(ws, row: int) -> tuple[str, int] | None:
    """(stripped label, indent depth) for column A of `row`, or None for a
    blank row. Depth is the leading-space count halved (this template
    indents two spaces per nesting level)."""
    raw = ws.cell(row=row, column=1).value
    if raw is None or not str(raw).strip():
        return None
    text = str(raw)
    return text.strip(), (len(text) - len(text.lstrip(" "))) // 2


def find_section_row(ws, label: str, max_row: int) -> int | None:
    """First row whose column-A label matches `label` exactly (any depth) --
    used to locate a named section header, e.g. "Income Statement"."""
    for row in range(1, max_row + 1):
        parsed = _row_label_depth(ws, row)
        if parsed and parsed[0].lower() == label.lower():
            return row
    return None


def find_label_in_range(ws, start_row: int, end_row: int, label: str, depth: int) -> int | None:
    """First row in [start_row, end_row) whose column-A label matches
    `label` exactly at exactly `depth` -- the depth filter is what
    disambiguates a repeated label like "Net Income" (the statement's own
    line vs. a nested "Adjusted Results" variant a few rows below it)."""
    for row in range(start_row, end_row):
        parsed = _row_label_depth(ws, row)
        if parsed and parsed[1] == depth and parsed[0].lower() == label.lower():
            return row
    return None


def sum_child_rows(ws, header_row: int, header_depth: int, col: str, max_row: int) -> float | None:
    """Sum numeric values in `col` across the rows nested under header_row
    (depth > header_depth), stopping at the next row at or above that
    depth. Used when a subtotal row (e.g. "Changes in Working Capital") is
    left blank in the export but its components underneath it aren't."""
    col_idx = column_index_from_string(col)
    total = None
    for row in range(header_row + 1, max_row + 1):
        parsed = _row_label_depth(ws, row)
        if parsed is None:
            continue
        if parsed[1] <= header_depth:
            break
        value = ws.cell(row=row, column=col_idx).value
        if isinstance(value, (int, float)):
            total = (total or 0) + value
    return total


def load_historicals_by_label(ws) -> list[dict]:
    """Fallback historicals loader for a label-only export (see module note
    above): resolves every concept by (section, label, depth) via
    REQUIRED_LABEL_MAP/OPTIONAL_LABEL_MAP instead of a Bloomberg field
    code. Raises the same way load_historicals() does when something
    load-bearing can't be found."""
    max_row = ws.max_row
    section_rows = {name: find_section_row(ws, name, max_row) for name in LABEL_SECTIONS}
    missing_sections = [name for name, row in section_rows.items() if row is None]
    if missing_sections:
        raise ValueError(
            f"Could not find section(s) {missing_sections} in this workbook, and no "
            "Bloomberg field codes were found either -- this layout isn't one this "
            "loader recognizes."
        )
    income_row, bs_row, cf_row = (section_rows[name] for name in LABEL_SECTIONS)
    bounds = {
        "Income Statement": (income_row, bs_row),
        "Condensed Balance Sheet": (bs_row, cf_row),
        "Condensed Cash Flow Statement": (cf_row, max_row + 1),
    }

    def resolve(section: str, label: str, depth: int) -> int | None:
        start, end = bounds[section]
        return find_label_in_range(ws, start + 1, end, label, depth)

    required_rows = {concept: resolve(*spec) for concept, spec in REQUIRED_LABEL_MAP.items()}
    missing = [concept for concept, row in required_rows.items() if row is None]
    if missing:
        raise ValueError(
            f"Could not find line item(s) for {missing} under their expected section "
            f"in this label-only export. Tried: {[REQUIRED_LABEL_MAP[c] for c in missing]}."
        )
    optional_rows = {concept: resolve(*spec) for concept, spec in OPTIONAL_LABEL_MAP.items()}

    records = []
    for col in find_year_columns(ws):
        year = ws[f"{col}4"].value.year
        record = {"fiscal_year": year, "period_type": period_type_for_column(ws, col)}
        for concept, row in required_rows.items():
            record[concept] = _numeric_or_none(ws[f"{col}{row}"].value)
        for concept, row in optional_rows.items():
            record[concept] = _numeric_or_none(ws[f"{col}{row}"].value) if row else None

        if record.get("nwc_change") is None and optional_rows.get("nwc_change"):
            nwc_row = optional_rows["nwc_change"]
            nwc_depth = OPTIONAL_LABEL_MAP["nwc_change"][2]
            record["nwc_change"] = sum_child_rows(ws, nwc_row, nwc_depth, col, max_row)

        cash = record.pop("cash", None)
        record["net_debt"] = (record["long_term_debt"] or 0) - (cash or 0)
        if record["nwc_change"] is None:
            record["nwc_change"] = 0.0

        records.append(record)
    return records


# ------------------------------------------------------ Mode B: derive it --

def fade_to_terminal(known_values: list[float], total_years: int, terminal_value: float) -> list[float]:
    """First len(known_values) years use the known (consensus) values;
    remaining years fade linearly from the last known value to
    terminal_value. If there are more known values than total_years,
    only the first total_years are used."""
    faded = list(known_values[:total_years])
    remaining = total_years - len(faded)
    if remaining <= 0:
        return faded
    start = faded[-1] if faded else terminal_value
    for i in range(1, remaining + 1):
        faded.append(start + (terminal_value - start) * (i / remaining))
    return faded


def _mean(values: list[float]) -> float:
    values = [v for v in values if v is not None]
    return sum(values) / len(values) if values else 0.0


def derive_scenario_assumptions(historicals: list[dict]) -> dict[str, list[dict]]:
    """Formulaically derive Bear/Base/Bull scenario assumptions purely from
    `historicals` — no hand-built DCF tab required. See module docstring.
    (Terminal exit multiple is computed separately by the caller — it
    needs current market data, which this function doesn't take.)
    """
    rows = sorted(historicals, key=lambda h: h["fiscal_year"])
    actual = [h for h in rows if h["period_type"] == "actual"]
    estimate = [h for h in rows if h["period_type"] == "estimate"]

    consensus_growth = []
    prior_revenue = actual[-1]["revenue"] if actual else None
    for h in estimate:
        if prior_revenue:
            consensus_growth.append(h["revenue"] / prior_revenue - 1)
        prior_revenue = h["revenue"]

    all_known = actual + estimate
    ebit_margin = _mean([h["ebit"] / h["revenue"] for h in all_known if h["revenue"]])
    tax_rate = _mean([h["tax_expense"] / h["pretax_income"] for h in all_known if h["pretax_income"] and h["pretax_income"] > 0])
    da_pct = _mean([h["da"] / h["revenue"] for h in all_known if h["revenue"]])
    capex_pct = _mean([abs(h["capex"]) / h["revenue"] for h in all_known if h["revenue"] and h["capex"]])
    # NWC change as a share of revenue (rather than of ΔRevenue, which would
    # need a clean consecutive-year series) — 0 when the field was absent
    # for every year (load_historicals() already defaults missing NWC data
    # to 0 per year).
    nwc_pct = _mean([h["nwc_change"] / h["revenue"] for h in all_known if h["revenue"]])

    base_growth = fade_to_terminal(consensus_growth, TOTAL_FORECAST_YEARS, TERMINAL_GROWTH_RATE)

    def make_years(growth: list[float], margin: float) -> list[dict]:
        return [
            {
                "fiscal_year": FIRST_FORECAST_YEAR + i,
                "revenue_growth": g,
                "ebit_margin": margin,
                "tax_rate": tax_rate,
                "da_pct_revenue": da_pct,
                "capex_pct_revenue": capex_pct,
                "nwc_pct_delta_revenue": nwc_pct,
            }
            for i, g in enumerate(growth)
        ]

    return {
        "base": make_years(base_growth, ebit_margin),
        "bear": make_years([g * BEAR_GROWTH_MULTIPLIER for g in base_growth], ebit_margin + BEAR_MARGIN_DELTA),
        "bull": make_years([g * BULL_GROWTH_MULTIPLIER for g in base_growth], ebit_margin + BULL_MARGIN_DELTA),
    }


def load_workbook_data(xlsx_path: Path, market_data: dict | None = None) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    mp_ws = find_multiple_periods_sheet(wb)
    if mp_ws is None:
        raise ValueError(
            f"This workbook has no 'Multiple Periods' sheet (found {wb.sheetnames}). "
            "That sheet is required in every mode — it's the historicals/consensus data."
        )

    # A field-code MODL export has Bloomberg codes (IS_COMP_SALES etc.) in
    # column B; a raw export saved straight from the on-screen grid doesn't
    # -- see load_historicals_by_label()'s module note.
    mp_field_rows = find_all_field_rows(mp_ws)
    has_field_codes = any(code in mp_field_rows for aliases in REQUIRED_FIELD_ALIASES.values() for code in aliases)
    load_mp_historicals = load_historicals if has_field_codes else load_historicals_by_label

    has_dcf_wacc = "DCF" in wb.sheetnames and "WACC" in wb.sheetnames

    if has_dcf_wacc:
        # Mode A: a hand-built DCF/WACC tab exists — extract it verbatim, unchanged.
        dcf_ws = wb["DCF"]
        wacc_ws = wb["WACC"]
        scenarios = {}
        terminal_multiples = {}
        for scenario, header_row in SCENARIO_BLOCKS.items():
            assumptions, exit_multiple = load_scenario_block(dcf_ws, header_row)
            scenarios[scenario] = assumptions
            terminal_multiples[scenario] = exit_multiple
        return {
            "ticker": parse_ticker(mp_ws),
            "name": parse_company_name_from_title(dcf_ws["A1"].value or ""),
            "as_of_date": parse_as_of_date(dcf_ws["A2"].value or ""),
            "data_source": "modl_tabs",
            "historicals": load_mp_historicals(mp_ws),
            "wacc_inputs": load_wacc_inputs(wacc_ws),
            "scenarios": scenarios,
            "terminal_multiples": terminal_multiples,
            "trading_comps": load_trading_comps(dcf_ws),
        }

    # Mode B: no hand-built tabs — the normal shape of a raw MODL export.
    # Compute everything from historicals; market/CAPM inputs must be
    # supplied explicitly (never fabricated).
    missing_market_data = [k for k in REQUIRED_MARKET_DATA_KEYS if not market_data or k not in market_data]
    if missing_market_data:
        raise ValueError(
            f"This workbook only has 'Multiple Periods' (no DCF/WACC tabs), so it needs "
            f"market_data supplied explicitly for {missing_market_data} — none of that is "
            "in any Bloomberg financials export (it's market/pricing data, not company "
            "data), so it can't be derived and won't be guessed."
        )

    historicals = load_mp_historicals(mp_ws)
    scenarios = derive_scenario_assumptions(historicals)
    latest_actual = max((h for h in historicals if h["period_type"] == "actual"), key=lambda h: h["fiscal_year"])

    market_cap = market_data["stock_price"] * latest_actual["shares_diluted"]
    current_ev_ebitda = (market_cap + latest_actual["net_debt"]) / latest_actual["ebitda_adjusted"]
    terminal_multiples = {
        "base": current_ev_ebitda,
        "bear": current_ev_ebitda * BEAR_MULTIPLE_MULTIPLIER,
        "bull": current_ev_ebitda * BULL_MULTIPLE_MULTIPLIER,
    }

    return {
        "ticker": parse_ticker(mp_ws),
        "name": parse_company_name_from_title(mp_ws["A1"].value or ""),
        "as_of_date": None,
        "data_source": "derived",
        "historicals": historicals,
        "wacc_inputs": {
            "risk_free_rate": market_data["risk_free_rate"],
            "beta": market_data["beta"],
            "equity_risk_premium": market_data["equity_risk_premium"],
            "stock_price": market_data["stock_price"],
            "shares_outstanding": latest_actual["shares_diluted"],
            "source": market_data.get("source") or "Informado manualmente no upload",
        },
        "scenarios": scenarios,
        "terminal_multiples": terminal_multiples,
        "trading_comps": [],  # no market-implied multiples without a DCF tab; P/E and
                               # PEG target-price methods degrade gracefully (see app.py).
    }


def _ensure_column(conn: sqlite3.Connection, table: str, column: str, coltype: str) -> None:
    """Lightweight migration: add a column to an existing table if it's not
    there yet, so a database created before this column existed doesn't
    break on the next load."""
    existing = {row[1] for row in conn.execute(f"PRAGMA table_info({table})")}
    if column not in existing:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {coltype}")


def migrate_schema(conn: sqlite3.Connection, schema_path: Path) -> None:
    """Bring a database's schema up to date: create any missing tables
    (schema.sql uses CREATE TABLE IF NOT EXISTS) and add any columns
    added since the database was first created. Safe to call on every
    connection, not just on load -- a database that was populated before
    a column existed (e.g. the committed data/valuation.db before
    data_source was added) would otherwise crash the first time
    something reads that column, even without a new upload."""
    conn.executescript(schema_path.read_text())
    conn.execute("PRAGMA foreign_keys = ON")
    _ensure_column(conn, "historicals", "shares_diluted", "REAL")
    _ensure_column(conn, "companies", "data_source", "TEXT NOT NULL DEFAULT 'modl_tabs'")
    _ensure_column(conn, "wacc_inputs", "source", "TEXT")


def write_to_db(db_path: Path, schema_path: Path, data: dict) -> None:
    conn = sqlite3.connect(db_path)
    try:
        migrate_schema(conn, schema_path)

        existing = conn.execute(
            "SELECT id FROM companies WHERE ticker = ?", (data["ticker"],)
        ).fetchone()
        if existing:
            company_id = existing[0]
            for table in (
                "historicals", "wacc_inputs", "scenario_assumptions",
                "terminal_assumptions", "trading_comps",
            ):
                conn.execute(f"DELETE FROM {table} WHERE company_id = ?", (company_id,))
            conn.execute(
                "UPDATE companies SET name = ?, as_of_date = ?, data_source = ? WHERE id = ?",
                (data["name"], data["as_of_date"], data["data_source"], company_id),
            )
        else:
            cur = conn.execute(
                "INSERT INTO companies (ticker, name, as_of_date, data_source) VALUES (?, ?, ?, ?)",
                (data["ticker"], data["name"], data["as_of_date"], data["data_source"]),
            )
            company_id = cur.lastrowid

        for row in data["historicals"]:
            conn.execute(
                """INSERT INTO historicals
                   (company_id, fiscal_year, period_type, revenue, ebit, ebitda_adjusted,
                    da, interest_expense, pretax_income, tax_expense, net_income,
                    eps_diluted_adjusted, long_term_debt, net_debt, capex, nwc_change,
                    shares_diluted)
                   VALUES (:company_id, :fiscal_year, :period_type, :revenue, :ebit,
                    :ebitda_adjusted, :da, :interest_expense, :pretax_income, :tax_expense,
                    :net_income, :eps_diluted_adjusted, :long_term_debt, :net_debt, :capex,
                    :nwc_change, :shares_diluted)""",
                {**row, "company_id": company_id, "shares_diluted": row.get("shares_diluted")},
            )

        conn.execute(
            """INSERT INTO wacc_inputs
               (company_id, risk_free_rate, beta, equity_risk_premium, stock_price, shares_outstanding, source)
               VALUES (:company_id, :risk_free_rate, :beta, :equity_risk_premium, :stock_price, :shares_outstanding, :source)""",
            {**data["wacc_inputs"], "company_id": company_id, "source": data["wacc_inputs"].get("source")},
        )

        for scenario, assumptions in data["scenarios"].items():
            for row in assumptions:
                conn.execute(
                    """INSERT INTO scenario_assumptions
                       (company_id, scenario, fiscal_year, revenue_growth, ebit_margin, tax_rate,
                        da_pct_revenue, capex_pct_revenue, nwc_pct_delta_revenue)
                       VALUES (:company_id, :scenario, :fiscal_year, :revenue_growth, :ebit_margin,
                        :tax_rate, :da_pct_revenue, :capex_pct_revenue, :nwc_pct_delta_revenue)""",
                    {**row, "company_id": company_id, "scenario": scenario},
                )
            conn.execute(
                """INSERT INTO terminal_assumptions (company_id, scenario, exit_multiple_ebitda)
                   VALUES (?, ?, ?)""",
                (company_id, scenario, data["terminal_multiples"][scenario]),
            )

        for row in data["trading_comps"]:
            conn.execute(
                """INSERT INTO trading_comps (company_id, metric, fiscal_year, value)
                   VALUES (:company_id, :metric, :fiscal_year, :value)""",
                {**row, "company_id": company_id},
            )

        conn.commit()
    finally:
        conn.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx_path", type=Path, help="Path to the Bloomberg MODL .xlsx export")
    parser.add_argument("db_path", type=Path, help="Path to the SQLite database to create/update")
    parser.add_argument("--risk-free-rate", type=float, help="10Y Treasury yield, e.g. 0.0477 (Mode B only)")
    parser.add_argument("--beta", type=float, help="5Y monthly beta (Mode B only)")
    parser.add_argument("--erp", type=float, help="Equity risk premium, e.g. 0.0445 (Mode B only)")
    parser.add_argument("--stock-price", type=float, help="Current stock price (Mode B only)")
    args = parser.parse_args()

    market_data = None
    if args.risk_free_rate is not None or args.beta is not None or args.erp is not None or args.stock_price is not None:
        market_data = {
            "risk_free_rate": args.risk_free_rate,
            "beta": args.beta,
            "equity_risk_premium": args.erp,
            "stock_price": args.stock_price,
        }

    schema_path = Path(__file__).parent / "schema.sql"
    data = load_workbook_data(args.xlsx_path, market_data=market_data)
    args.db_path.parent.mkdir(parents=True, exist_ok=True)
    write_to_db(args.db_path, schema_path, data)
    print(f"Loaded {data['ticker']} ({data['name']}) into {args.db_path}")


if __name__ == "__main__":
    main()
