"""Regression tests for the valuation calculator.

Run with: pytest

Pins down the numbers this project was validated against (the source
Bloomberg MODL file's $388.54 base-case price, and the three-way
sensitivity-table mismatch it also contains), plus the edge cases found
during a full-project review: explicit_years bounds and target_year
data-coverage gaps that could otherwise crash app.py.
"""

from __future__ import annotations

import shutil
import sqlite3
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from data.load_from_modl import (
    _numeric_or_none,
    derive_scenario_assumptions,
    load_workbook_data,
    write_to_db,
)
from dcf_engine import get_company_id, get_wacc, run_dcf
from sensitivity import (
    sensitivity_beta_risk_free,
    sensitivity_growth_margin,
    sensitivity_wacc_exit_multiple,
)
from target_price import football_field, from_ev_ebitda, from_pe, from_peg

DB_PATH = Path(__file__).parent.parent / "data" / "valuation.db"
TICKER = "APP US"

CLONE_COLUMNS = {
    "historicals": "fiscal_year,period_type,revenue,ebit,ebitda_adjusted,da,"
                   "interest_expense,pretax_income,tax_expense,net_income,"
                   "eps_diluted_adjusted,long_term_debt,net_debt,capex,nwc_change",
    "wacc_inputs": "risk_free_rate,beta,equity_risk_premium,stock_price,shares_outstanding",
    "scenario_assumptions": "scenario,fiscal_year,revenue_growth,ebit_margin,tax_rate,"
                             "da_pct_revenue,capex_pct_revenue,nwc_pct_delta_revenue",
    "terminal_assumptions": "scenario,exit_multiple_ebitda",
    "trading_comps": "metric,fiscal_year,value",
}


@pytest.fixture
def conn():
    connection = sqlite3.connect(DB_PATH)
    yield connection
    connection.close()


@pytest.fixture
def multi_company_conn(tmp_path):
    """A temp copy of valuation.db with a second, cloned company under a
    different ticker, to prove company_id scoping doesn't bleed across
    companies (the schema is designed to be multi-company but this was
    never actually tested with more than one before this review)."""
    db_copy = tmp_path / "multi.db"
    shutil.copy(DB_PATH, db_copy)
    connection = sqlite3.connect(db_copy)
    cur = connection.cursor()
    cur.execute(
        "INSERT INTO companies (ticker, name, currency, units, as_of_date) "
        "SELECT 'CLONE US', 'Clone Corp', currency, units, as_of_date "
        "FROM companies WHERE ticker = ?", (TICKER,),
    )
    new_id = cur.lastrowid
    old_id = get_company_id(connection, TICKER)
    for table, cols in CLONE_COLUMNS.items():
        cur.execute(
            f"INSERT INTO {table} (company_id,{cols}) SELECT ?,{cols} FROM {table} WHERE company_id = ?",
            (new_id, old_id),
        )
    connection.commit()
    yield connection
    connection.close()


@pytest.fixture
def clone_conn_factory(tmp_path):
    """Factory for a temp copy of valuation.db with a cloned 'BROKEN US'
    company whose latest actual historicals row can be tweaked (e.g. zero
    debt), used to exercise get_wacc()'s division-by-zero guards without
    touching real AppLovin data."""

    def _make(**historicals_overrides):
        db_copy = tmp_path / f"broken_{len(historicals_overrides)}_{'-'.join(historicals_overrides)}.db"
        shutil.copy(DB_PATH, db_copy)
        connection = sqlite3.connect(db_copy)
        cur = connection.cursor()
        cur.execute(
            "INSERT INTO companies (ticker, name, currency, units, as_of_date) "
            "SELECT 'BROKEN US', 'Broken Corp', currency, units, as_of_date "
            "FROM companies WHERE ticker = ?", (TICKER,),
        )
        new_id = cur.lastrowid
        old_id = get_company_id(connection, TICKER)
        for table, cols in CLONE_COLUMNS.items():
            cur.execute(
                f"INSERT INTO {table} (company_id,{cols}) SELECT ?,{cols} FROM {table} WHERE company_id = ?",
                (new_id, old_id),
            )
        if historicals_overrides:
            set_clause = ", ".join(f"{col} = ?" for col in historicals_overrides)
            cur.execute(
                f"""UPDATE historicals SET {set_clause}
                    WHERE company_id = ? AND period_type = 'actual'
                    AND fiscal_year = (SELECT MAX(fiscal_year) FROM historicals
                                       WHERE company_id = ? AND period_type = 'actual')""",
                (*historicals_overrides.values(), new_id, new_id),
            )
        connection.commit()
        return connection

    return _make


# --------------------------------------------------------- DCF baseline --

def test_dcf_base_case_matches_source_file(conn):
    """C87 in the source Bloomberg MODL workbook, cell-by-cell validated."""
    result = run_dcf(conn, TICKER, "base")
    assert result["price_per_share"] == pytest.approx(388.54471895377645)


def test_dcf_thirteen_year_alternative(conn):
    """The one internally-consistent variant of the source file's own
    sensitivity tables (D110) — see README 'Segunda descoberta'."""
    result = run_dcf(conn, TICKER, "base", explicit_years=13)
    assert result["price_per_share"] == pytest.approx(412.1852731339656)


@pytest.mark.parametrize("bad_years", [0, -1, 14, 20])
def test_dcf_rejects_out_of_range_explicit_years(conn, bad_years):
    """explicit_years=0 used to IndexError; negative values silently sliced
    the projection list backwards and returned a wrong-but-plausible price."""
    with pytest.raises(ValueError):
        run_dcf(conn, TICKER, "base", explicit_years=bad_years)


def test_dcf_rejects_unknown_scenario(conn):
    with pytest.raises(ValueError):
        run_dcf(conn, TICKER, "mid")


def test_dcf_rejects_unknown_ticker(conn):
    with pytest.raises(ValueError):
        run_dcf(conn, "FAKE US", "base")


# --------------------------------------------------------- sensitivity --

def test_sensitivity_wacc_multiple_center_cell_matches_dcf(conn):
    dcf_price = run_dcf(conn, TICKER, "base")["price_per_share"]
    table = sensitivity_wacc_exit_multiple(conn, TICKER, "base", explicit_years=5)
    center = table["grid"][2][2]  # zero WACC delta, zero multiple delta
    assert center == pytest.approx(dcf_price)


def test_sensitivity_growth_margin_center_cell_matches_dcf(conn):
    dcf_price = run_dcf(conn, TICKER, "base")["price_per_share"]
    table = sensitivity_growth_margin(conn, TICKER, "base", explicit_years=5)
    center = table["grid"][2][2]
    assert center == pytest.approx(dcf_price)


def test_sensitivity_reproduces_source_files_own_thirteen_year_table(conn):
    """B108/F112 in the source file's Growth x Margin sensitivity table —
    the one internally-consistent one (see README)."""
    table = sensitivity_growth_margin(conn, TICKER, "base", explicit_years=13)
    assert table["grid"][0][0] == pytest.approx(204.29703188199804)
    assert table["grid"][-1][-1] == pytest.approx(865.4601924527196)


# ------------------------------------------------------------- multiples --

def test_from_ev_ebitda_matches_known_value(conn):
    result = from_ev_ebitda(conn, TICKER, scenario="base", target_year=2027)
    assert result["target_price"] == pytest.approx(344.146541586244)


def test_from_pe_raises_for_year_without_trading_comps(conn):
    """FY2026E/FY2028E/FY2029E have no market-implied P/E in the source
    file's TRADING MULTIPLES block (only FY2027E/FY2030E do). Found while
    reviewing why app.py's target_year dropdown could crash the page."""
    with pytest.raises(ValueError):
        from_pe(conn, TICKER, target_year=2026)


@pytest.mark.parametrize("year", [2027, 2030])
def test_from_pe_works_for_covered_years(conn, year):
    result = from_pe(conn, TICKER, target_year=year)
    assert result["target_price"] > 0


def test_from_peg_does_not_depend_on_trading_comps(conn):
    """Unlike P/E, PEG derives its own multiple from EPS CAGR, so it works
    for every consensus year (2026-2030) even though trading_comps only
    covers 2027/2030."""
    for year in range(2026, 2031):
        result = from_peg(conn, TICKER, target_year=year)
        assert result["target_price"] > 0


# -------------------------------------------------- multi-company scoping --

def test_multi_company_results_do_not_bleed(multi_company_conn):
    original = run_dcf(multi_company_conn, TICKER, "base")["price_per_share"]
    clone = run_dcf(multi_company_conn, "CLONE US", "base")["price_per_share"]
    assert original == pytest.approx(clone)  # CLONE US is a byte-for-byte copy


# --------------------------------------------------- get_wacc() guards --

def test_get_wacc_rejects_zero_debt(clone_conn_factory):
    """AppLovin carries debt so this never fires today, but the schema is
    meant to hold other companies too, including debt-free ones."""
    conn = clone_conn_factory(long_term_debt=0)
    with pytest.raises(ValueError, match="debt-free"):
        get_wacc(conn, get_company_id(conn, "BROKEN US"))


def test_get_wacc_rejects_zero_pretax_income(clone_conn_factory):
    conn = clone_conn_factory(pretax_income=0)
    with pytest.raises(ValueError, match="effective tax rate"):
        get_wacc(conn, get_company_id(conn, "BROKEN US"))


def test_get_wacc_rejects_zero_enterprise_value(clone_conn_factory):
    # net_debt == -market_cap (317.76 * 335.94) makes enterprise value 0.
    conn = clone_conn_factory(net_debt=-106748.2944)
    with pytest.raises(ValueError, match="enterprise value"):
        get_wacc(conn, get_company_id(conn, "BROKEN US"))


# ---------------------------------------------------- from_peg() guards --

def test_from_peg_rejects_negative_base_eps(conn):
    """FY2022 EPS (-0.52) would otherwise raise a negative base to a
    fractional power — Python silently returns a complex number instead
    of erroring, which then breaks formatting far from the real cause."""
    with pytest.raises(ValueError, match="must be positive"):
        from_peg(conn, TICKER, target_year=2026, base_year=2022)


# ------------------------------------------------- Beta x Risk-Free grid --

def test_sensitivity_beta_risk_free_center_cell_matches_dcf(conn):
    dcf_price = run_dcf(conn, TICKER, "base")["price_per_share"]
    table = sensitivity_beta_risk_free(conn, TICKER, "base", explicit_years=5)
    center = table["grid"][2][2]  # zero risk-free delta, zero beta delta
    assert center == pytest.approx(dcf_price)


# ------------------------------------------- load_from_modl() sheet check --

def test_load_workbook_data_rejects_missing_multiple_periods_sheet(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Some Other Sheet"
    path = tmp_path / "no_multiple_periods.xlsx"
    wb.save(path)

    with pytest.raises(ValueError, match="Multiple Periods"):
        load_workbook_data(path)


def test_load_workbook_data_mode_b_requires_market_data(tmp_path):
    """A workbook with only 'Multiple Periods' (no DCF/WACC tabs) is the
    normal shape of a raw Bloomberg export (confirmed against a real second
    company, Blackstone) -- Mode B. CAPM/market inputs (risk-free rate,
    beta, ERP, stock price) are never in a financials export, so
    load_workbook_data() must refuse to guess them and say exactly what's
    missing rather than silently defaulting or crashing with a bare
    openpyxl KeyError."""
    wb = openpyxl.Workbook()
    wb.active.title = "Multiple Periods"
    path = tmp_path / "historicals_only.xlsx"
    wb.save(path)

    with pytest.raises(ValueError, match="market_data"):
        load_workbook_data(path)


# --------------------------------------- Mode B: derive_scenario_assumptions --

FIRST_ACTUAL_YEAR = 2021
YEARS = list(range(FIRST_ACTUAL_YEAR, FIRST_ACTUAL_YEAR + 10))  # 4 actual + 6 estimate
ACTUAL_YEARS = YEARS[:4]
REVENUES = [1000 * (1.1 ** i) for i in range(10)]


def _synthetic_historicals() -> list[dict]:
    """10 years of made-up-but-internally-consistent historicals: 20% EBIT
    margin, 5% D&A, 2% interest expense, 21% tax rate, all as a share of
    revenue, and constant 10% revenue growth throughout -- deterministic
    enough to assert exact derived numbers against."""
    records = []
    for i, year in enumerate(YEARS):
        revenue = REVENUES[i]
        ebit = revenue * 0.2
        da = revenue * 0.05
        interest_expense = revenue * 0.02
        pretax_income = ebit - interest_expense
        tax_expense = pretax_income * 0.21
        records.append({
            "fiscal_year": year,
            "period_type": "actual" if year in ACTUAL_YEARS else "estimate",
            "revenue": revenue,
            "ebit": ebit,
            "ebitda_adjusted": ebit + da,
            "da": da,
            "interest_expense": interest_expense,
            "pretax_income": pretax_income,
            "tax_expense": tax_expense,
            "net_income": pretax_income - tax_expense,
            "eps_diluted_adjusted": (pretax_income - tax_expense) / 100,
            "long_term_debt": 500.0,
            "net_debt": 500.0,
            "capex": -revenue * 0.05,
            "nwc_change": 0.0,
            "shares_diluted": 100.0,
        })
    return records


def test_derive_scenario_assumptions_growth_matches_consensus_then_fades():
    scenarios = derive_scenario_assumptions(_synthetic_historicals())
    base = scenarios["base"]
    assert len(base) == 13
    # 6 estimate years (2025-2030) all grew 10% over the prior year -- the
    # known/consensus part of the fade should reproduce that exactly.
    for year_data in base[:6]:
        assert year_data["revenue_growth"] == pytest.approx(0.1)
    # The last of the 7 faded years lands exactly on the terminal rate.
    assert base[-1]["revenue_growth"] == pytest.approx(0.03)
    # Trailing-average margins/rates, held flat across every year.
    for year_data in base:
        assert year_data["ebit_margin"] == pytest.approx(0.2)
        assert year_data["tax_rate"] == pytest.approx(0.21)
        assert year_data["da_pct_revenue"] == pytest.approx(0.05)
        assert year_data["capex_pct_revenue"] == pytest.approx(0.05)
        assert year_data["nwc_pct_delta_revenue"] == pytest.approx(0.0)


def test_derive_scenario_assumptions_bear_bull_envelope():
    scenarios = derive_scenario_assumptions(_synthetic_historicals())
    base, bear, bull = scenarios["base"], scenarios["bear"], scenarios["bull"]
    assert bear[0]["revenue_growth"] == pytest.approx(base[0]["revenue_growth"] * 0.7)
    assert bull[0]["revenue_growth"] == pytest.approx(base[0]["revenue_growth"] * 1.3)
    assert bear[0]["ebit_margin"] == pytest.approx(base[0]["ebit_margin"] - 0.03)
    assert bull[0]["ebit_margin"] == pytest.approx(base[0]["ebit_margin"] + 0.02)


# ------------------------------------- Mode B: load_workbook_data end-to-end --

def _build_mode_b_workbook(path: Path) -> None:
    """A minimal synthetic 'Multiple Periods'-only workbook (no DCF/WACC
    tabs), shaped like a real raw Bloomberg MODL export, using the same
    field-code layout _synthetic_historicals() mirrors as plain dicts."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Multiple Periods"
    ws["A1"] = "Synthetic Test Corp (TEST)"
    ws["A2"] = "TEST US Equity"

    field_rows = {
        "IS_COMP_SALES": "revenue",
        "IS_COMPARABLE_EBIT": "ebit",
        "IS_COMPARABLE_EBITDA": "ebitda_adjusted",
        "CF_DEPR_AMORT": "da",
        "IS_NET_INTEREST_EXPENSE": "interest_expense",
        "PRETAX_INC": "pretax_income",
        "IS_INC_TAX_EXP": "tax_expense",
        "IS_COMP_NET_INCOME_GAAP": "net_income",
        "IS_COMP_EPS_ADJUSTED_OLD": "eps_diluted_adjusted",
        "CB_BS_LT_BORROWING": "long_term_debt",
        "CB_CF_PURCHASES_OF_PPE": "capex",
        "IS_SH_FOR_DILUTED_EPS": "shares_diluted",
    }
    cols = ["E", "F", "G", "H", "I", "J", "K", "L", "M", "N"]
    for row, (code, key) in enumerate(field_rows.items(), start=10):
        ws[f"B{row}"] = code
        for col, record in zip(cols, _synthetic_historicals()):
            ws[f"{col}{row}"] = record[key]

    for col, record, year in zip(cols, _synthetic_historicals(), YEARS):
        ws[f"{col}3"] = f"{year} A (Rep)" if record["period_type"] == "actual" else f"{year} A (Fwd)"
        ws[f"{col}4"] = date(year, 12, 31)

    wb.save(path)


def test_load_workbook_data_mode_b_computes_from_historicals(tmp_path):
    path = tmp_path / "mode_b.xlsx"
    _build_mode_b_workbook(path)
    market_data = {"risk_free_rate": 0.04, "beta": 1.2, "equity_risk_premium": 0.05, "stock_price": 50.0}

    data = load_workbook_data(path, market_data=market_data)

    assert data["ticker"] == "TEST US"
    assert data["name"] == "Synthetic Test Corp"
    assert data["data_source"] == "derived"
    assert data["trading_comps"] == []
    assert data["wacc_inputs"] == {
        "risk_free_rate": 0.04,
        "beta": 1.2,
        "equity_risk_premium": 0.05,
        "stock_price": 50.0,
        "shares_outstanding": 100.0,
        "source": "Informado manualmente no upload",
    }
    # Latest actual year (2024): revenue 1331, EBITDA 1331*0.25=332.75,
    # net_debt 500, market_cap 50*100=5000.
    expected_ev_ebitda = (5000 + 500) / 332.75
    assert data["terminal_multiples"]["base"] == pytest.approx(expected_ev_ebitda)
    assert data["terminal_multiples"]["bear"] == pytest.approx(expected_ev_ebitda * 0.85)
    assert data["terminal_multiples"]["bull"] == pytest.approx(expected_ev_ebitda * 1.15)
    assert len(data["scenarios"]["base"]) == 13


def test_mode_b_company_round_trips_through_db_and_football_field(tmp_path):
    """End-to-end: a derived (Mode B) company writes to a fresh db and
    dcf_engine/target_price can price it -- P/E is unavailable (no
    trading_comps) but DCF, EV/EBITDA and PEG still work, mirroring what
    football_field() showed for the real Blackstone file."""
    xlsx_path = tmp_path / "mode_b.xlsx"
    _build_mode_b_workbook(xlsx_path)
    market_data = {"risk_free_rate": 0.04, "beta": 1.2, "equity_risk_premium": 0.05, "stock_price": 50.0}
    data = load_workbook_data(xlsx_path, market_data=market_data)

    db_path = tmp_path / "mode_b.db"
    schema_path = Path(__file__).parent.parent / "data" / "schema.sql"
    write_to_db(db_path, schema_path, data)

    conn = sqlite3.connect(db_path)
    try:
        results = football_field(conn, "TEST US", "base")
        methods = {r["method"] for r in results}
        assert methods == {"DCF", "EV/EBITDA", "PEG"}
        with pytest.raises(ValueError):
            from_pe(conn, "TEST US", target_year=2027)
    finally:
        conn.close()


# ------------------------------- label-only export (no field codes at all) --
# A raw "Company Financial (Multiple Periods)" export saved straight from
# Bloomberg's on-screen grid -- rather than the field-code-driven MODL
# template -- has no BDH-style field-code column (column B holds the first
# year's data, not a code) and its tab can be named anything (e.g. the
# spreadsheet app's own default "Sheet1"), found by matching "(Multiple
# Periods)" in the sheet's own A1 title text instead. Found by testing
# against a real third-party file (a Bloomberg export of Alphabet/GOOGL
# saved without any manual renaming) -- built here as a minimal synthetic
# workbook, never committing that real file.

LABEL_YEARS = [2024, 2025, 2026, 2027]  # 2 actual + 2 estimate
LABEL_ACTUAL_YEARS = LABEL_YEARS[:2]
LABEL_COLS = ["B", "C", "D", "E"]  # starts at B, not E -- no field-code column to skip


def _build_label_only_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"  # deliberately NOT "Multiple Periods" -- title-only detection
    ws["A1"] = "Test Corp- Company Financial (Multiple Periods)"
    ws["A2"] = "TEST US Equity    Periodicity:A    Currency:USD"

    for col, year in zip(LABEL_COLS, LABEL_YEARS):
        ws[f"{col}3"] = f"{year} A (Rep)" if year in LABEL_ACTUAL_YEARS else f"{year} A (Fwd)"
        ws[f"{col}4"] = date(year, 12, 31)

    # (row, indent depth, label)
    rows = [
        (10, 1, "Income Statement"),
        (11, 1, "Total Revenue"),
        (12, 1, "Operating Income"),
        (13, 2, "EBITDA"),
        (14, 1, "Depreciation & Amortization"),
        (15, 2, "Interest Expense"),
        (16, 1, "Pre-Tax Income"),
        (17, 1, "Income Tax Expense"),
        (18, 1, "Net Income"),  # the one that should be picked (depth 1)
        (19, 2, "Net Income"),  # decoy at a different depth -- must NOT be picked
        (20, 1, "Diluted Weighted Avg. Shares"),
        (21, 1, "Diluted EPS"),  # GAAP decoy -- must NOT be picked
        (22, 2, "Diluted EPS"),  # adjusted -- the one that should be picked
        (30, 1, "Condensed Balance Sheet"),
        (31, 3, "Cash Cash Equivalents and Short term Investments"),
        (32, 3, "Long-Term Debt"),
        (40, 1, "Condensed Cash Flow Statement"),
        (41, 2, "Capital Expenditures"),
        (42, 2, "Changes in Working Capital"),  # left blank -- sum_child_rows() must fill it
        (43, 3, "Accounts Receivables"),
        (44, 3, "Accounts Payable"),
    ]
    for row, depth, label in rows:
        ws.cell(row=row, column=1, value="  " * depth + label)

    revenues = [1000 * (1.1 ** i) for i in range(4)]
    for col, revenue in zip(LABEL_COLS, revenues):
        ebit = revenue * 0.2
        da = revenue * 0.05
        pretax = ebit * 0.9
        tax = pretax * 0.21
        net_income = pretax - tax
        shares = 100.0
        ws[f"{col}11"] = revenue
        ws[f"{col}12"] = ebit
        ws[f"{col}13"] = revenue * 0.28  # ebitda_adjusted, distinct from ebit+da on purpose
        ws[f"{col}14"] = da
        ws[f"{col}15"] = revenue * 0.01  # gross interest expense
        ws[f"{col}16"] = pretax
        ws[f"{col}17"] = tax
        ws[f"{col}18"] = net_income
        ws[f"{col}19"] = net_income * 1.05  # decoy value -- must not be read
        ws[f"{col}20"] = shares
        ws[f"{col}21"] = net_income / shares * 0.90  # GAAP decoy -- must not be read
        ws[f"{col}22"] = net_income / shares * 0.95  # adjusted -- must be read
        ws[f"{col}31"] = revenue * 0.1  # cash
        ws[f"{col}32"] = 500.0  # long-term debt
        ws[f"{col}41"] = -revenue * 0.05  # capex
        # row 42 ("Changes in Working Capital") intentionally left blank
        ws[f"{col}43"] = -50.0
        ws[f"{col}44"] = 20.0

    wb.save(path)


def test_numeric_or_none_normalizes_blank_string():
    assert _numeric_or_none("") is None
    assert _numeric_or_none("   ") is None
    assert _numeric_or_none(0) == 0
    assert _numeric_or_none(12.5) == 12.5
    assert _numeric_or_none(None) is None


def test_load_workbook_data_label_only_export(tmp_path):
    path = tmp_path / "label_only.xlsx"
    _build_label_only_workbook(path)
    market_data = {"risk_free_rate": 0.04, "beta": 1.1, "equity_risk_premium": 0.045, "stock_price": 40.0}

    data = load_workbook_data(path, market_data=market_data)

    assert data["ticker"] == "TEST US"
    assert data["name"] == "Test Corp"
    assert data["data_source"] == "derived"
    assert len(data["historicals"]) == 4
    assert [h["fiscal_year"] for h in data["historicals"]] == LABEL_YEARS
    assert [h["period_type"] for h in data["historicals"]] == ["actual", "actual", "estimate", "estimate"]

    fy2024 = data["historicals"][0]
    revenue = 1000.0
    ebit = revenue * 0.2
    pretax = ebit * 0.9
    tax = pretax * 0.21
    net_income = pretax - tax
    # The depth-1 "Net Income" is read, not the depth-2 decoy right below it.
    assert fy2024["net_income"] == pytest.approx(net_income)
    # The depth-2 "Diluted EPS" (adjusted) is read, not the depth-1 GAAP decoy.
    assert fy2024["eps_diluted_adjusted"] == pytest.approx(net_income / 100.0 * 0.95)
    # Gross interest expense (row 15), not a net-of-interest-income figure.
    assert fy2024["interest_expense"] == pytest.approx(revenue * 0.01)
    # "Changes in Working Capital" itself was left blank -- summed from its
    # two child rows (Accounts Receivables -50, Accounts Payable +20).
    assert fy2024["nwc_change"] == pytest.approx(-30.0)
    assert fy2024["long_term_debt"] == pytest.approx(500.0)
    assert fy2024["net_debt"] == pytest.approx(500.0 - revenue * 0.1)


def test_load_workbook_data_label_only_export_rejects_missing_section(tmp_path):
    """A workbook with no Bloomberg field codes AND missing one of the three
    expected statement sections isn't a layout this loader recognizes --
    it should say so, not silently return wrong numbers or a cryptic
    AttributeError."""
    path = tmp_path / "no_cash_flow_section.xlsx"
    _build_label_only_workbook(path)

    wb = openpyxl.load_workbook(path)
    ws = wb["Sheet1"]
    ws["A40"] = None  # remove the "Condensed Cash Flow Statement" header
    wb.save(path)

    market_data = {"risk_free_rate": 0.04, "beta": 1.1, "equity_risk_premium": 0.045, "stock_price": 40.0}
    with pytest.raises(ValueError, match="Condensed Cash Flow Statement"):
        load_workbook_data(path, market_data=market_data)
